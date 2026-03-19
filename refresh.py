#!/usr/bin/env python3
"""
refresh.py — Semantic Scholar paper refresh for UIV India Deeptech Database.

Usage:
    python refresh.py "IIT Madras"
    python refresh.py "CSIR" --limit 20
    python refresh.py "IIT Bombay" --domain "quantum computing" --limit 10
    python refresh.py --list            # List known institutes

The script:
1. Queries Semantic Scholar Graph API (free, no API key required)
2. Deduplicates against existing rows in the .xlsx file
3. Appends new papers and updates the Master_Index.xlsx summary count
"""

import sys
import os
import time
import json
import argparse
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ── Constants ────────────────────────────────────────────────────────────────

NAVY  = "0D1B4B"
GOLD  = "C9A84C"
WHITE = "FFFFFF"

S2_BASE = "https://api.semanticscholar.org/graph/v1/paper/search"
S2_FIELDS = "title,authors,year,externalIds,abstract,venue,publicationTypes"
RATE_LIMIT_DELAY = 1.0   # seconds between API calls (polite rate)
MAX_RETRIES = 3

# ── Institute registry ───────────────────────────────────────────────────────

INSTITUTE_MAP = {
    "iit madras":   {"file": "IIT_Madras.xlsx",   "queries": ["IIT Madras deeptech", "IIT Madras research"]},
    "iitm":         {"file": "IIT_Madras.xlsx",   "queries": ["IIT Madras deeptech", "IIT Madras research"]},
    "csir":         {"file": "CSIR.xlsx",          "queries": ["CSIR India research", "CSIR laboratory India"]},
    "iit bombay":   {"file": "IIT_Bombay.xlsx",   "queries": ["IIT Bombay deeptech", "IIT Bombay research"]},
    "iitb":         {"file": "IIT_Bombay.xlsx",   "queries": ["IIT Bombay deeptech", "IIT Bombay research"]},
    "iit delhi":    {"file": "IIT_Delhi.xlsx",    "queries": ["IIT Delhi deeptech", "IIT Delhi research"]},
    "iitd":         {"file": "IIT_Delhi.xlsx",    "queries": ["IIT Delhi deeptech", "IIT Delhi research"]},
    "iit hyderabad":{"file": "IIT_Hyderabad.xlsx","queries": ["IIT Hyderabad deeptech", "IIT Hyderabad research"]},
    "iith":         {"file": "IIT_Hyderabad.xlsx","queries": ["IIT Hyderabad deeptech", "IIT Hyderabad research"]},
    "iit kharagpur":{"file": "IIT_Kharagpur.xlsx","queries": ["IIT Kharagpur deeptech", "IIT Kharagpur research"]},
    "iitkgp":       {"file": "IIT_Kharagpur.xlsx","queries": ["IIT Kharagpur deeptech"]},
    "iisc":         {"file": "IISc.xlsx",          "queries": ["IISc Bangalore research", "Indian Institute of Science deeptech"]},
    "drdo":         {"file": "DRDO.xlsx",          "queries": ["DRDO India research", "Defence Research Development Organisation"]},
    "isro":         {"file": "ISRO.xlsx",          "queries": ["ISRO India space research", "Indian Space Research Organisation"]},
    "nrdc":         {"file": "NRDC.xlsx",          "queries": ["NRDC India technology", "National Research Development Corporation"]},
    "barc":         {"file": "BARC.xlsx",          "queries": ["BARC India nuclear research", "Bhabha Atomic Research Centre"]},
}

# ── Style helpers ─────────────────────────────────────────────────────────────

def _thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _navy_fill():
    return PatternFill("solid", fgColor=NAVY)

def _alt_fill(row):
    return PatternFill("solid", fgColor="EEF0F8") if row % 2 == 0 else PatternFill("solid", fgColor=WHITE)

def _style_new_row(ws, row, ncols):
    """Apply light-green highlight to freshly appended rows."""
    new_fill = PatternFill("solid", fgColor="D4EDDA")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = new_fill
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = _thin_border()

# ── Semantic Scholar API ──────────────────────────────────────────────────────

def fetch_papers(query: str, limit: int = 10) -> list[dict]:
    """Query Semantic Scholar and return list of paper dicts."""
    params = {
        "query": query,
        "limit": min(limit, 100),
        "fields": S2_FIELDS,
    }
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.get(S2_BASE, params=params, timeout=15)
            if resp.status_code == 200:
                data = resp.json()
                return data.get("data", [])
            elif resp.status_code == 429:
                wait = 2 ** attempt
                print(f"  Rate limited. Waiting {wait}s...")
                time.sleep(wait)
            else:
                print(f"  API error {resp.status_code}: {resp.text[:200]}")
                return []
        except requests.RequestException as e:
            print(f"  Network error (attempt {attempt}/{MAX_RETRIES}): {e}")
            time.sleep(2 ** attempt)
    return []


def normalise_paper(raw: dict, domain: str = "") -> dict:
    """Convert Semantic Scholar raw result to our schema."""
    authors = ", ".join(a.get("name", "") for a in raw.get("authors", [])[:5])
    if len(raw.get("authors", [])) > 5:
        authors += " et al."
    ids = raw.get("externalIds", {}) or {}
    doi = ids.get("DOI", "")
    doi_url = f"https://doi.org/{doi}" if doi else ids.get("ArXiv", "")
    if ids.get("ArXiv") and not doi_url:
        doi_url = f"https://arxiv.org/abs/{ids['ArXiv']}"

    abstract = raw.get("abstract") or ""
    if len(abstract) > 300:
        abstract = abstract[:297] + "..."

    # Simple heuristic UIV relevance
    title_lower = (raw.get("title") or "").lower()
    relevance = "Medium"
    high_kw = ["defence", "defense", "military", "security", "quantum", "nuclear",
                "autonomous", "satellite", "missile", "cyber", "drone", "radar",
                "hypersonic", "surveil", "biodefense"]
    low_kw = ["survey", "review of", "history of", "pedagogy", "education"]
    if any(k in title_lower for k in high_kw):
        relevance = "High — Strategic / Defence"
    elif any(k in title_lower for k in low_kw):
        relevance = "Low"

    return {
        "title":         raw.get("title", ""),
        "authors":       authors,
        "year":          raw.get("year", ""),
        "journal":       raw.get("venue", ""),
        "domain":        domain or "General",
        "doi":           doi_url,
        "abstract":      abstract,
        "uiv_relevance": relevance,
    }

# ── xlsx helpers ─────────────────────────────────────────────────────────────

PAPER_COLS = ["Title", "Authors", "Year", "Journal/Conference", "Domain",
              "DOI/URL", "Abstract Summary", "UIV Relevance"]

def get_existing_titles(ws) -> set[str]:
    """Return set of lowercase titles already in the Research Papers sheet."""
    titles = set()
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            titles.add(str(row[0]).strip().lower())
    return titles


def ensure_papers_sheet(wb) -> openpyxl.worksheet.worksheet.Worksheet:
    """Return Research Papers sheet, creating it if missing."""
    if "Research Papers" in wb.sheetnames:
        return wb["Research Papers"]
    # Create minimal sheet
    ws = wb.create_sheet("Research Papers")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(PAPER_COLS))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(PAPER_COLS))
    t = ws.cell(row=1, column=1, value="Research Papers")
    t.fill = _navy_fill()
    t.font = Font(name="Calibri", bold=True, color=GOLD, size=16)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36
    for c, h in enumerate(PAPER_COLS, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill = _navy_fill()
        cell.font = Font(name="Calibri", bold=True, color=WHITE, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()
    ws.row_dimensions[3].height = 22
    widths = [40, 25, 8, 30, 20, 30, 50, 20]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return ws


def append_papers(ws, new_papers: list[dict]) -> int:
    """Append papers to sheet, return count added."""
    existing = get_existing_titles(ws)
    next_row = ws.max_row + 1
    # Ensure we're past the header (row 3)
    if next_row < 4:
        next_row = 4

    added = 0
    for p in new_papers:
        if p["title"].strip().lower() in existing:
            continue
        row = next_row + added
        ws.row_dimensions[row].height = 40
        vals = [p["title"], p["authors"], p["year"], p["journal"],
                p["domain"], p["doi"], p["abstract"], p["uiv_relevance"]]
        for c, v in enumerate(vals, 1):
            ws.cell(row=row, column=c, value=v)
        _style_new_row(ws, row, len(PAPER_COLS))
        existing.add(p["title"].strip().lower())
        added += 1
    return added


def update_master_index(institute_file: str, new_paper_count: int):
    """Update paper count in Master_Index.xlsx for the given institute file."""
    master_path = "Master_Index.xlsx"
    if not os.path.exists(master_path):
        print(f"  Master_Index.xlsx not found — skipping update.")
        return
    wb = openpyxl.load_workbook(master_path)
    if "Institute Directory" not in wb.sheetnames:
        return
    ws = wb["Institute Directory"]
    for row in ws.iter_rows(min_row=4):
        if row[3].value == institute_file:   # column D = File
            current = row[7].value or 0      # column H = Papers
            row[7].value = int(current) + new_paper_count
            ts_cell = row[11]                # column L = Last Refreshed
            ts_cell.value = datetime.now().strftime("%Y-%m-%d")
            break
    # Also update Summary Stats sheet
    if "Summary Stats" in wb.sheetnames:
        ws2 = wb["Summary Stats"]
        for row in ws2.iter_rows(min_row=3):
            if row[0].value == "Total Research Papers":
                row[1].value = (row[1].value or 0) + new_paper_count
                break
    wb.save(master_path)
    print(f"  Updated Master_Index.xlsx (+{new_paper_count} papers for {institute_file})")


# ── Main ──────────────────────────────────────────────────────────────────────

def resolve_institute(name: str) -> tuple[str, list[str]] | None:
    """Return (xlsx_filename, query_list) or None."""
    key = name.strip().lower()
    if key in INSTITUTE_MAP:
        info = INSTITUTE_MAP[key]
        return info["file"], info["queries"]
    # Fuzzy: check if key is a substring
    for k, v in INSTITUTE_MAP.items():
        if key in k or k in key:
            return v["file"], v["queries"]
    return None


def main():
    parser = argparse.ArgumentParser(
        description="Refresh Research Papers sheet from Semantic Scholar API"
    )
    parser.add_argument("institute", nargs="?", help="Institute name or code (e.g. 'IIT Madras', 'CSIR')")
    parser.add_argument("--domain", default="", help="Optional domain context to append to query")
    parser.add_argument("--limit", type=int, default=15, help="Max papers per query (default 15)")
    parser.add_argument("--list", action="store_true", help="List known institutes and exit")
    parser.add_argument("--dry-run", action="store_true", help="Fetch but don't write to xlsx")
    args = parser.parse_args()

    if args.list:
        print("\nKnown institutes:")
        seen_files = {}
        for k, v in INSTITUTE_MAP.items():
            f = v["file"]
            if f not in seen_files:
                seen_files[f] = k
                print(f"  {k:20s} → {f}")
        return

    if not args.institute:
        parser.print_help()
        sys.exit(1)

    result = resolve_institute(args.institute)
    if not result:
        print(f"Unknown institute: '{args.institute}'. Use --list to see options.")
        sys.exit(1)

    xlsx_file, base_queries = result
    queries = [f"{q} {args.domain}".strip() for q in base_queries]

    print(f"\n{'='*60}")
    print(f"  Institute : {args.institute}")
    print(f"  File      : {xlsx_file}")
    print(f"  Queries   : {queries}")
    print(f"  Limit     : {args.limit} per query")
    print(f"{'='*60}\n")

    # Collect papers
    all_raw = []
    for q in queries:
        print(f"  Querying Semantic Scholar: '{q}'...")
        results = fetch_papers(q, limit=args.limit)
        print(f"    → {len(results)} results")
        all_raw.extend(results)
        time.sleep(RATE_LIMIT_DELAY)

    if not all_raw:
        print("  No papers fetched. Check network or try a different query.")
        return

    domain_label = args.domain or "General / Mixed"
    normalised = [normalise_paper(r, domain_label) for r in all_raw]

    print(f"\n  Total fetched (before dedup): {len(normalised)}")

    if args.dry_run:
        print("\n  [DRY RUN] Sample papers:")
        for p in normalised[:3]:
            print(f"    • {p['year']} — {p['title'][:70]}")
        return

    # Load or create xlsx
    if os.path.exists(xlsx_file):
        wb = openpyxl.load_workbook(xlsx_file)
        print(f"  Loaded existing: {xlsx_file}")
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "Overview"
        print(f"  Created new: {xlsx_file}")

    ws = ensure_papers_sheet(wb)
    added = append_papers(ws, normalised)

    if added > 0:
        wb.save(xlsx_file)
        print(f"  Saved {xlsx_file} (+{added} new papers appended, highlighted green)")
        update_master_index(xlsx_file, added)
    else:
        print("  No new papers to add (all already present).")

    print(f"\nDone. {added}/{len(normalised)} new papers added.")


if __name__ == "__main__":
    main()
